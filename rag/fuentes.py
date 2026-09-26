import re
import unicodedata
from collections import Counter
from pathlib import Path

import pymupdf as fitz
import yaml
from docx import Document
from pptx import Presentation

EXTENSIONES = {".pdf", ".docx", ".pptx"}

TIPOS = [
    ("evaluacion", r"exam|quiz|parcial"),
    ("ejercicios", r"ejercicio|exercise|taller|practica|complementaria|enunciado|as \d"),
]


def leer_pdf(ruta):
    paginas = []
    with fitz.open(ruta) as doc:
        for i, pagina in enumerate(doc):
            bloques = [b[4] for b in pagina.get_text("blocks") if b[6] == 0]
            paginas.append((i + 1, [re.sub(r"\s*\n\s*", " ", b).strip() for b in bloques]))
    return paginas


def leer_docx(ruta):
    doc = Document(ruta)
    parrafos = [p.text.strip() for p in doc.paragraphs]
    for tabla in doc.tables:
        for fila in tabla.rows:
            parrafos.append(" | ".join(c.text.strip() for c in fila.cells))
    return [(1, parrafos)]


def textos_forma(forma):
    if forma.shape_type == 6:
        return [t for f in forma.shapes for t in textos_forma(f)]
    if forma.has_table:
        return [" | ".join(c.text.strip() for c in fila.cells) for fila in forma.table.rows]
    return [forma.text_frame.text.strip()] if forma.has_text_frame else []


def leer_pptx(ruta):
    return [
        (i + 1, [t for forma in d.shapes for t in textos_forma(forma)])
        for i, d in enumerate(Presentation(ruta).slides)
    ]


LECTORES = {".pdf": leer_pdf, ".docx": leer_docx, ".pptx": leer_pptx}


def quitar_repetidos(paginas):
    if len(paginas) < 3:
        return paginas
    conteo = Counter(p for _, parrafos in paginas for p in set(parrafos) if len(p) < 120)
    repetidos = {p for p, n in conteo.items() if n >= 0.6 * len(paginas)}
    return [(n, [p for p in parrafos if p not in repetidos]) for n, parrafos in paginas]


def arreglar_tildes(texto):
    texto = texto.replace("\u00b4\u0131", "\u00ed").replace("\u0131\u00b4", "\u00ed")
    texto = re.sub(r"\u00b4\s?([aeiouAEIOU])", lambda m: m.group(1) + "\u0301", texto)
    texto = re.sub(r"\u02dc\s?([nN])", lambda m: m.group(1) + "\u0303", texto)
    return unicodedata.normalize("NFC", texto)


def limpiar(parrafos):
    salida = []
    for p in parrafos:
        p = arreglar_tildes(p)
        p = re.sub(r"(\w)- (\w)", r"\1\2", p)
        p = re.sub(r"[ \t]+", " ", p).strip()
        if p and not re.fullmatch(r"\d{1,3}|página \d+.*|page \d+.*", p, re.I):
            salida.append(p)
    return salida


def leer(ruta):
    paginas = LECTORES[ruta.suffix.lower()](ruta)
    return [(n, limpiar(parrafos)) for n, parrafos in quitar_repetidos(paginas)]


def tipo_material(nombre):
    for tipo, patron in TIPOS:
        if re.search(patron, nombre, re.I):
            return tipo
    return "teoria"


def idioma(texto):
    palabras = re.findall(r"\b\w+\b", texto.lower())
    es = sum(w in {"de", "la", "el", "que", "los", "una", "para"} for w in palabras)
    en = sum(w in {"the", "of", "and", "is", "that", "for", "with"} for w in palabras)
    return "en" if en > es else "es"


def tutorias(raiz):
    for ficha in sorted(Path(raiz).glob("*/tutoria.yaml")):
        datos = yaml.safe_load(ficha.read_text(encoding="utf-8")) or {}
        carpeta = ficha.parent
        archivos = sorted(
            (f for f in carpeta.rglob("*") if f.suffix.lower() in EXTENSIONES),
            key=lambda f: f.stat().st_size,
        )
        yield {
            "tutoria": carpeta.name,
            "estudiante": datos.get("estudiante") or "",
            "curso": datos.get("curso") or "",
            "temas": ", ".join(datos.get("temas") or []),
        }, archivos
